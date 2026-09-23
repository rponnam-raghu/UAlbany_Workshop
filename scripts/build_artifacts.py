"""Build the fictional intake PDF and Word FAQ used by the workshop sample pack."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "data" / "samples"


def pdf(path: Path, deadline: str, details: str) -> None:
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="WorkshopTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, leading=25, textColor=colors.HexColor("#163c2c"), spaceAfter=12, alignment=TA_LEFT))
    styles.add(ParagraphStyle(name="WorkshopHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#163c2c"), spaceBefore=12, spaceAfter=5))
    styles.add(ParagraphStyle(name="WorkshopBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=10.3, leading=15, spaceAfter=8))
    doc = SimpleDocTemplate(str(path), pagesize=letter, rightMargin=.72 * inch, leftMargin=.72 * inch, topMargin=.7 * inch, bottomMargin=.7 * inch)
    story = [Paragraph("Spring 2027 Intake Guide", styles["WorkshopTitle"]), Paragraph("FICTIONAL WORKSHOP UNIVERSITY", styles["WorkshopHeading"]), Paragraph("This document is fictional workshop material. It is not an official University at Albany deadline or admissions guide.", styles["WorkshopBody"]), Paragraph("Application deadline", styles["WorkshopHeading"]), Paragraph(f"The Spring 2027 application closes on <b>{deadline}</b> at 11:59 p.m. Eastern Time.", styles["WorkshopBody"]), Paragraph(details, styles["WorkshopBody"]), Paragraph("Required materials", styles["WorkshopHeading"])]
    table = Table([["Material", "What to prepare"], ["Application form", "Completed fictional application form"], ["Academic record", "Unofficial transcript"], ["Statement", "Short statement of academic goals"], ["Reference", "Contact information for one academic reference"]], colWidths=[1.65 * inch, 4.95 * inch])
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dce9df")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#163c2c")), ("GRID", (0, 0), (-1, -1), .5, colors.HexColor("#c8d4cb")), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, 1), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), 9.5), ("LEADING", (0, 0), (-1, -1), 12), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
    story += [table, Spacer(1, 12), Paragraph("Questions about exceptions, late materials, or eligibility should go to a human advisor at advising@workshop.example. The helpdesk can explain this guide but cannot approve an exception.", styles["WorkshopBody"])]
    doc.build(story)


def docx(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.7)
    section.left_margin = Inches(.8)
    section.right_margin = Inches(.8)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    title_style_properties = styles["Title"]._element.get_or_add_pPr()
    title_border = title_style_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_style_properties.remove(title_border)
    for name in ("Heading 1", "Heading 2"):
        styles[name].font.name = "Aptos Display"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Advising FAQ")
    document.add_paragraph("Fictional workshop material for practicing retrieval, citations, and reliable answers. It is not official University at Albany policy.")
    document.add_heading("Changing a course plan", level=1)
    document.add_paragraph("Students may discuss a course-plan change with an advisor before registering. The helpdesk can explain prerequisites and catalog offerings, but a human advisor handles exceptions and approval decisions.")
    document.add_heading("Course eligibility", level=1)
    document.add_paragraph("Eligibility depends on completed courses and the current catalog. A course offered in a particular semester may still be unavailable to a student who has not completed its prerequisites.")
    document.add_heading("Common questions", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Question"
    table.rows[0].cells[1].text = "Planning answer"
    rows = [("Can the helpdesk approve an exception?", "No. A human advisor or registrar handles approvals."), ("Can I rely on a past answer after a document changes?", "Ask again. The current uploaded documents are the evidence."), ("What should I include when I contact advising?", "The topic, planning term, and relevant course code. Do not email personal transcripts to this workshop application.")]
    for question, answer in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = question, answer
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    if row is table.rows[0]:
                        run.bold = True
    document.add_heading("Contact", level=1)
    document.add_paragraph("Email fictional advising questions to advising@workshop.example. The helpdesk explains reference material; it does not approve enrollment or certify graduation.")
    document.save(path)


if __name__ == "__main__":
    SAMPLES.mkdir(parents=True, exist_ok=True)
    pdf(SAMPLES / "intake_guide.pdf", "November 15, 2026", "Applications submitted after that time are reviewed only if an advisor confirms that late materials are being accepted.")
    docx(SAMPLES / "advising_faq.docx")
