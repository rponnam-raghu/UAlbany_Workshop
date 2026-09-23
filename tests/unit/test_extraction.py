from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from academic_advisor.knowledge.extraction import UploadError, chunk, extract


def make_pdf() -> bytes:
    return (Path("data/samples/intake_guide.pdf")).read_bytes()


def make_docx() -> bytes:
    document = Document()
    document.add_heading("FAQ", level=1)
    document.add_paragraph("A human advisor handles approvals.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Question"
    table.rows[0].cells[1].text = "Answer"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_pdf_preserves_page_reference():
    upload = extract("guide.pdf", make_pdf())
    assert upload.sections[0].page == 1
    assert "November 15" in upload.sections[0].text


def test_docx_preserves_table_reference():
    upload = extract("faq.docx", make_docx())
    assert any("Table 1" in section.location for section in upload.sections)
    assert any("Question" in section.text for section in upload.sections)


@pytest.mark.parametrize("filename,content", [("bad.exe", b"x"), ("empty.txt", b""), ("binary.txt", b"x\x00\x01")])
def test_invalid_uploads_are_explained(filename, content):
    with pytest.raises(UploadError):
        extract(filename, content)


def test_scanned_like_pdf_is_rejected():
    stream = BytesIO()
    PdfWriter().write(stream)
    with pytest.raises(UploadError, match="OCR"):
        extract("scan.pdf", stream.getvalue())


def test_chunk_keeps_source_locations():
    upload = extract("notes.md", b"# Heading\n\n" + b"word " * 500)
    passages = chunk(upload.sections, size=100, overlap=20)
    assert passages
    assert all(p.location == "Heading" for p in passages)
